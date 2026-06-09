import os
import pandas as pd
import logging
import glob
import csv
from datetime import datetime
import openpyxl
import math
import sys
import matplotlib.pyplot as plt
import seaborn as sns
import numpy as np
from matplotlib.colors import LinearSegmentedColormap
import re
from matplotlib.table import Table
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
import warnings
warnings.filterwarnings("ignore")
from collections import defaultdict
from PySide6.QtCore import QThread, Signal
import time
import traceback

class ProcessingWorkertable(QThread):
    job_starting = Signal(int)
    job_completed = Signal(int)
    job_failed = Signal(int,str)
    def __init__(self,path_n,vehicle_varient,vehicle_name,battery_name,Tyre,V_type,gear_ratio):
        super().__init__()
        self.V_variant = vehicle_varient
        self.V_name = vehicle_name
        self.battery = battery_name
        self.TYRE = Tyre
        self.path_new = os.path.abspath(os.path.dirname(path_n[-1]))
        self.Vehicle_type = V_type
        self.g_ratio = gear_ratio
    def run(self):
        self.job_starting.emit(0)
        self.table_png()

    def table_png(self):
        try:
            battery = self.battery
            firm=gvw=payload=0
            if self.V_variant in ["V1","V2","V3"]:
                name = f"{self.V_variant}-{self.V_name}"
            else:
                name = self.V_name
            file_path = self.path_new
            df = pd.read_csv(os.path.join(file_path, "Final_summary.csv"))
            if self.Vehicle_type == "4W":
                if self.V_variant == "1750":
                    payload = "1100Kg"
                    gvw = "2708"
                    firm = "0"
                elif self.V_variant == "1250":
                    firm = "v1.1.0.1.1\n config-3 F/W-6"
                    payload = "1250Kg"
                    if self.V_name == "S2":
                        gvw = "2718"
                    elif self.V_name == "S3":
                        gvw = "2678"
                    else:
                        gvw = "2698"
                elif self.V_variant == "LR200":
                    firm = "v1.9.0.9.0\n config-1 F/W-6"
                    payload = "1300Kg"
                    gvw = "2980"
                elif self.V_variant == "V1":
                    firm = "v1.0.0.6.3\n config-3 F/W-4"
                    if self.V_name == "7009":
                        payload = "2000Kg"
                        gvw = "3233"
                    else:
                        payload = "800Kg"
                        gvw = "2033"
                elif self.V_variant == "V2":
                    firm = "v1.0.0.3.0\n config-2 F/W-5"
                    if self.V_name == "7019":
                        payload = "2000Kg"
                        gvw = "3348"
                    else:
                        payload = "800Kg"
                        gvw = "2148"
                elif self.V_variant == "V3":
                    firm = "v1.2.0.1.1\n config-5 F/W-7"  
                    if self.V_name == "7008":
                        payload = "2500Kg"
                        gvw = "3868"
                    else:
                        payload = "800Kg"
                        gvw = "2168" 
                Payload = payload
                Peak_torque = df["Motor Peak Torque(Nm)"].iloc[0]
                start_imbalance = self.fmt(df["Start Imbalance"].iloc[0])
                end_imbalance = self.fmt(df["End Imbalance"].iloc[0])
                peak_imbalance = self.fmt(df["Peak Imbalance"].iloc[0])
                max_cell_temp = df["Peak Cell Temp"].iloc[0]
                Mode_firmware = firm
            elif self.Vehicle_type == "3W":#-------<<< yet to complete >>>-----
                print("it went in 3W checkpoint")
                start_motor_temp = self.fmt(df["Motor Initial Temp"].iloc[0])
                peak_ARMS = self.fmt(df["Controller Peak ARMS"].iloc[0])
                avg_ARMS = self.fmt(df["Controller Avg ARMS"].iloc[0])
                whkm = self.fmt(df["Wh/KM"].iloc[0])
                avg_rpm = self.fmt(df["Avg. Motor RPM"].iloc[0])
                peak_rpm = self.fmt(df["Motor Peak RPM"].iloc[0])
                gvw = "nill"
                mcu = "nill"
                vehicle_dict = {
                    "SR_HL": ("GTake","1540"), 
                    "TR_HL": ("GTake","1540"),
                    "XR_HL": ("GTake","1560"),
                    "SR_HC": ("GTake","1010"),
                    "TR_HC": ("Pegasus","1010"),#GTake,
                    "XR_HC": ("GTake","1030"),
                    "SR_HR": ("GTake","900"),
                    "TR_HR": ("Pegasus","900"),#GTake,
                    "XR_HR": ("GTake","920"),
                    "SR_NEO": ("GTake","800"),
                    "TR_NEO": ("GTake","800"),
                    "XR_NEO": ("GTake","800"),
                    "SR_HR_2.0": ("GTake","760"),
                    "TR_HR_2.0": ("Pegasus","770"),
                    "XR_HR_2.0": ("GTake","800"),#,Pegasus
                    "SR_HL_2.0": ("Pegasus","1350"), 
                    "TR_HL_2.0": ("Pegasus","1350"),
                    "XR_HL_2.0": ("Pegasus","1350")
                }
                for n, values in vehicle_dict.items():
                    if n == self.V_variant:
                        mcu = values[0]
                        gvw = values[1]
                if self.V_name == "PT6":
                    mcu = "GTake"
                elif self.V_name == "PT1":
                    mcu = "GTake"
                elif self.V_name =="P4":
                    mcu = "Pegasus"
                    
                        
                

            time_str = df["Start Time"].iloc[0]
            time_obj = datetime.strptime(time_str, "%H:%M:%S")
            Time = time_obj.strftime("%I:%M %p")#C
            Date = df["Start Date"].iloc[0] # second row, first column
            Battery_pack = battery#C
            Equivalent_model = name#C
            GVW = gvw #vehicle based
            distance = self.fmt(df["Odometer Trip Range( )"].iloc[0])#C
            Top_speed = self.fmt(df["Max speed"].iloc[0])#C
            avg_speed = self.fmt(df["Avg Speed (>= 5Kmph)"].iloc[0])#C
            Max_motor_temp = df["Motor Peak temp (Max operating Temp: 145'C)"].iloc[0]#C
            regen = self.fmt(df["Regen Percent Recovery"].iloc[0])#C
            avg_current = self.fmt(df["Avg Battery Current"].iloc[0])#C
            peak_current = self.fmt(df["Battery Peak Current"].iloc[0])#C
            Tyre = self.TYRE#C
            
            

            if self.Vehicle_type == "4W":
                data = {
                "Date": f"{Date}\n{Time}",
                "Battery Pack \n (Rated/Designed)": Battery_pack,
                "Payload(kg)": Payload,
                "Equivalent model": Equivalent_model,
                "GVW(kg)": f"{GVW}Kg",
                "Range": f"{distance}Km",
                "Top Speed(kmph)": Top_speed,
                "(Avg. Speed)": avg_speed,
                "Peak Torque(Nm)": Peak_torque,
                "Max. Motor Temp.(℃)": Max_motor_temp,
                "Regen Percentage(%)": regen,
                "Start Imbalance(mV)": start_imbalance,
                "End Imbalance(mV)": end_imbalance,
                "Peak Imbalance(mV)": peak_imbalance,
                "Max. Cell Temp.(℃)": max_cell_temp,
                "Avg. Current(A)": avg_current,
                "Peak Current(A)": peak_current,
                "Mode (Firmware)": Mode_firmware,
                "TYRE": Tyre
                }
            elif self.Vehicle_type == "3W":
                MODE = self.fmt(df["Drive mode"].iloc[0])
                data = {
                "Date": f"{Date}\n{Time}",
                "Mode": MODE,
                "Battery Pack \n (Rated/Designed)": Battery_pack,
                "MCU / MOTOR": mcu,
                "Model": Equivalent_model,
                "Peak RPM": peak_rpm,
                "Avg RPM": avg_rpm,
                "Top Speed(kmph)": Top_speed,
                "(Avg. Speed)": avg_speed,
                "Range": f"{distance}Km",
                "Achieved WH/KM": whkm,
                "GVW(kg)": f"{GVW}Kg",
                "Regen Percentage(%)": regen,
                "Avg. Current(A)": avg_current,
                "Peak Current(A)": peak_current,
                "Avg ARMS": avg_ARMS,
                "Peak ARMS": peak_ARMS,
                "Motor Temp (start/ Peak)": f"{start_motor_temp} and {Max_motor_temp}",
                "Gear ratio":self.g_ratio,
                "TYRE": Tyre
                }
            self.save_styled_summary_table(data, self.path_new, name)
            self.create_summary_table(data,Equivalent_model,self.path_new)
            self.job_completed.emit(0)
        except Exception:
            error_string = traceback.format_exc()
            self.job_failed.emit(0,error_string)
            print(error_string)
    def save_styled_summary_table(self,data, path, name):
        # Convert dict to DataFrame
        summary_df = pd.DataFrame(list(data.items()), columns=["Parameter", "Value"])

        # Save raw DataFrame to Excel
        out_path = os.path.join(path, f"{name}_summary_table.xlsx")
        with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
            summary_df.to_excel(writer, index=False, sheet_name="Summary", startrow=0)

        # Load workbook
        wb = load_workbook(out_path)
        ws = wb["Summary"]

        # Define styles
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

        # Highlight colors mapping
        highlight_colors = {
            "Battery Pack": "00FF00",    # green
            "Equivalent model": "00FF00",
            "Range": "B6D7A8",          # light green
            "Top Speed": "9FC5E8",      # light blue
            "(Avg. Speed)": "9FC5E8",
            "Max. Motor Temp": "EA9999",# red
            "TYRE": "00FF00"
        }

        # Apply formatting to cells
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=2):
            param_cell, value_cell = row

            # Borders & alignment
            param_cell.border = thin_border
            value_cell.border = thin_border
            param_cell.alignment = left_align
            value_cell.alignment = center_align

            # Highlight specific rows
            for keyword, color in highlight_colors.items():
                if keyword in str(param_cell.value):
                    value_cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

        # Adjust column widths
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 25

        wb.save(out_path)
        print(f"Styled summary table saved as {out_path}")
    def fmt(self,x):
        if isinstance(x, (float, int)):
            return f"{x:.2f}"
        return x 
    def create_summary_table(self,data,model,path):
        fig, ax = plt.subplots(figsize=(5, 8))
        ax.set_axis_off()

        # Create table
        table = Table(ax, bbox=[0, 0, 1, 1])

        # Row height & column width
        n_rows = len(data)
        n_cols = 2
        col_width = 1.0 / n_cols
        base_height = 1.0 / n_rows
        # Define row heights (make first two rows taller)
        row_heights = [base_height * 1.8 if i < 2 else base_height for i in range(n_rows)]

        # Colors for categories
        highlight_colors = {
            "Battery Pack": "#00FF00",   # green
            "Equivalent model": "#00FF00",
            "Range": "#B6D7A8",
            "Top Speed": "#9FC5E8",
            "Avg. Speed": "#9FC5E8",
            "Max. Motor Temp.": "#EA9999",
            "TYRE": "#00FF00"
        }


        # Add cells
        for i, (key, value) in enumerate(data.items()):
            color = "white"
            for highlight in highlight_colors:
                if highlight in key:
                    color = highlight_colors[highlight]
            row_height = row_heights[i]
            # Add cells
            cell1 = table.add_cell(i, 0, col_width, row_height, text=key, loc="left", facecolor="white")
            cell2 = table.add_cell(i, 1, col_width, row_height, text=str(value), loc="center", facecolor=color)

            # Force larger font size
            cell1.get_text().set_fontsize(10)
            cell1.get_text().set_clip_on(False)
            cell2.get_text().set_fontsize(11)
            cell2.get_text().set_clip_on(False)

        # Remove table-level auto font scaling
        table.auto_set_font_size(False)

        # Increase row spacing instead of shrinking fonts
        table.scale(1, 2.0)

        ax.add_table(table)
        filename = f"{model}_table.png"
        save_location = os.path.join(path, filename)
        plt.savefig(save_location, dpi=300, bbox_inches="tight")
        print(f"Table saved as {filename}")
